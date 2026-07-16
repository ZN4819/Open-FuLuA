"""Deterministic OOXML renderer for the frozen complete-report master."""

from __future__ import annotations

import copy
import hashlib
import io
import mimetypes
import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from lxml import etree

from ..services.docx_generator.fields import BookmarkWriter
from ..services.docx_generator.images import add_section_images, build_figure_refs
from ..services.docx_generator.tables import add_assessment_table
from ..services.report_domain.errors import ReportDomainError
from ..services.report_templates.registry import report_template_registry
from ..services.template_profile import load_template_profile


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
XML = "http://www.w3.org/XML/1998/namespace"
NS = {"w": W, "r": R, "pr": REL, "m": M}
Q = lambda local: f"{{{W}}}{local}"

SEMANTIC_TAG_VALUES = {
    "report.identity.number": "report_number",
    "report.header.report_number": "report_number",
    "report.identity.date": "report_date",
    "report.organization.assessed_name": "assessed_name",
    "report.organization.assessment_name": "assessment_name",
    "report.declaration.assessment_name": "assessment_name",
    "report.header.assessment_name.1": "assessment_name",
    "report.header.assessment_name.2": "assessment_name",
    "report.cover.system_name": "system_name",
    "report.system.name": "system_name",
    "report.declaration.system_name": "system_name",
    "report.header.system_name.1": "system_name",
    "report.header.system_name.2": "system_name",
    "report.header.system_name.3": "system_name",
    "report.system.security_level": "security_level",
    "report.system.overview": "system_overview",
    "report.system.network_architecture": "network_architecture",
    "report.assessment.preparation_period": "preparation_period",
    "report.assessment.plan_period": "plan_period",
    "report.assessment.period": "assessment_period",
    "report.assessment.report_period": "report_period",
    "report.distribution.total_copies": "total_copies",
    "report.distribution.regulator_copies": "regulator_copies",
    "report.distribution.client_copies": "client_copies",
    "report.distribution.assessment_copies": "assessment_copies",
    "report.assessment.methods": "assessment_methods",
    "report.result.overall_score": "overall_score",
    "report.result.conclusion": "conclusion",
    "report.risk.high_risk_judgement": "high_risk_judgement",
}

LAYER_NAMES = (
    "物理和环境安全", "网络和通信安全", "设备和计算安全", "应用和数据安全",
    "管理制度", "人员管理", "建设运行", "应急处置",
)
PLACEHOLDER_RE = re.compile(r"【[^】]+】|(?<![A-Za-z])[xX]{2,}(?![A-Za-z])|待补充|填写说明|建议不超过\s*200\s*字|\{[^{}]+\}")


def _text(element: etree._Element) -> str:
    return "".join(element.xpath(".//w:t/text()", namespaces=NS)).strip()


def _paragraphs(root: etree._Element) -> list[etree._Element]:
    return list(root.xpath("//w:body/w:p", namespaces=NS))


def _find_paragraph(
    root: etree._Element,
    *,
    exact: str | None = None,
    startswith: str | None = None,
    contains: str | None = None,
    required: bool = True,
) -> etree._Element | None:
    matches = []
    for paragraph in _paragraphs(root):
        value = _text(paragraph)
        if (
            (exact is None or value == exact)
            and (startswith is None or value.startswith(startswith))
            and (contains is None or contains in value)
        ):
            matches.append(paragraph)
    if len(matches) == 1:
        return matches[0]
    if not required and not matches:
        return None
    raise ReportDomainError(
        "TEMPLATE_NARRATIVE_ANCHOR_INVALID",
        "母版正文白名单锚点缺失或不唯一。",
        status_code=500,
        details={"exact": exact, "startswith": startswith, "contains": contains, "count": len(matches)},
    )


def _new_run_properties(source: etree._Element | None, *, italic: bool = False) -> etree._Element | None:
    if source is None:
        return None
    existing = source.find(Q("rPr"))
    if existing is None:
        return None
    result = copy.deepcopy(existing)
    if not italic:
        for child in list(result):
            if child.tag in {Q("i"), Q("iCs")}:
                result.remove(child)
    return result


def _set_paragraph_text(
    paragraph: etree._Element,
    value: Any,
    *,
    italic: bool = False,
    first_line_chars: int | None = None,
) -> None:
    first_run = paragraph.find(Q("r"))
    rpr = _new_run_properties(first_run, italic=italic)
    ppr = paragraph.find(Q("pPr"))
    if first_line_chars is not None:
        if ppr is None:
            ppr = etree.Element(Q("pPr"))
            paragraph.insert(0, ppr)
        ind = ppr.find(Q("ind"))
        if ind is None:
            ind = etree.SubElement(ppr, Q("ind"))
        ind.set(Q("firstLineChars"), str(first_line_chars * 100))
    for child in list(paragraph):
        if child is not ppr:
            paragraph.remove(child)
    run = etree.SubElement(paragraph, Q("r"))
    if rpr is not None:
        run.append(rpr)
    text = etree.SubElement(run, Q("t"))
    rendered = str(value or "")
    if rendered.startswith((" ", "　")) or rendered.endswith(" "):
        text.set(f"{{{XML}}}space", "preserve")
    text.text = rendered


def _set_cell_text(cell: etree._Element, value: Any, *, bold: bool = False) -> None:
    tc_pr = cell.find(Q("tcPr"))
    first_p = cell.find(Q("p"))
    p_pr = copy.deepcopy(first_p.find(Q("pPr"))) if first_p is not None and first_p.find(Q("pPr")) is not None else None
    r_pr = None
    if first_p is not None:
        first_run = first_p.find(Q("r"))
        r_pr = _new_run_properties(first_run, italic=False)
    for child in list(cell):
        if child is not tc_pr:
            cell.remove(child)
    paragraph = etree.SubElement(cell, Q("p"))
    if p_pr is not None:
        paragraph.append(p_pr)
    run = etree.SubElement(paragraph, Q("r"))
    if r_pr is not None:
        run.append(r_pr)
    if bold:
        if r_pr is None:
            r_pr = etree.Element(Q("rPr"))
            run.insert(0, r_pr)
        etree.SubElement(r_pr, Q("b"))
    text = etree.SubElement(run, Q("t"))
    rendered = str(value if value not in (None, "") else "")
    if rendered.startswith((" ", "　")) or rendered.endswith(" "):
        text.set(f"{{{XML}}}space", "preserve")
    text.text = rendered


def _cells(row: etree._Element) -> list[etree._Element]:
    return list(row.xpath("./w:tc | ./w:sdt/w:sdtContent/w:tc", namespaces=NS))


def _set_sdt_text(root: etree._Element, tag: str, value: Any) -> None:
    matches = [
        sdt for sdt in root.xpath("//w:sdt", namespaces=NS)
        if sdt.xpath("string(w:sdtPr/w:tag/@w:val)", namespaces=NS) == tag
    ]
    if len(matches) != 1:
        raise ReportDomainError(
            "TEMPLATE_SCALAR_SLOT_INVALID", "母版语义字段槽位缺失或重复。", status_code=500,
            details={"tag": tag, "count": len(matches)},
        )
    texts = matches[0].xpath(".//w:sdtContent//w:t", namespaces=NS)
    if not texts:
        raise ReportDomainError(
            "TEMPLATE_SCALAR_SLOT_INVALID", "母版语义字段槽位没有可写文本节点。", status_code=500,
            details={"tag": tag},
        )
    texts[0].text = str(value if value not in (None, "") else "/")
    for node in texts[1:]:
        node.text = ""


def _table_by_anchor(root: etree._Element, number: int) -> etree._Element:
    name = f"rt_table_{number:03d}"
    starts = root.xpath(f"//w:bookmarkStart[@w:name='{name}']", namespaces=NS)
    if len(starts) != 1:
        raise ReportDomainError(
            "TEMPLATE_TABLE_ANCHOR_INVALID", "母版表格锚点缺失或重复。", status_code=500,
            details={"table_id": f"report_table_{number:03d}", "count": len(starts)},
        )
    tables = starts[0].xpath("ancestor::w:tbl[1]", namespaces=NS)
    if len(tables) != 1:
        raise ReportDomainError(
            "TEMPLATE_TABLE_ANCHOR_INVALID", "母版表格锚点未位于目标表格中。", status_code=500,
            details={"table_id": f"report_table_{number:03d}"},
        )
    return tables[0]


def _strip_row_merges(row: etree._Element) -> None:
    for merge in row.xpath(".//w:tcPr/w:vMerge", namespaces=NS):
        merge.getparent().remove(merge)


def _clone_row(template: etree._Element, values: list[Any]) -> etree._Element:
    row = copy.deepcopy(template)
    _strip_row_merges(row)
    cells = _cells(row)
    for index, cell in enumerate(cells):
        _set_cell_text(cell, values[index] if index < len(values) else "")
    return row


def _replace_rows(
    table: etree._Element,
    *,
    header_rows: int,
    source_rows: list[list[Any]],
    keep_summary: bool = False,
    summary_values: list[Any] | None = None,
) -> None:
    rows = list(table.xpath("./w:tr", namespaces=NS))
    if len(rows) <= header_rows:
        raise ReportDomainError("TEMPLATE_TABLE_ROW_CONTRACT_INVALID", "母版表格缺少模板数据行。", status_code=500)
    data_template = rows[header_rows]
    summary_template = rows[-1] if keep_summary else None
    for row in rows[header_rows:]:
        table.remove(row)
    materialized = source_rows or [["/"] * len(_cells(data_template))]
    for values in materialized:
        table.append(_clone_row(data_template, values))
    if summary_template is not None:
        table.append(_clone_row(summary_template, summary_values or ["/"] * len(_cells(summary_template))))


def _render_semantic_slots(parts: dict[str, bytes], scalars: dict[str, Any]) -> None:
    scalar_values = {**scalars, "assessment_methods": "访谈、文档审查、现场检查、配置检查、工具测试"}
    for name in list(parts):
        if name != "word/document.xml" and not re.fullmatch(r"word/(?:header|footer)\d+\.xml", name):
            continue
        root = etree.fromstring(parts[name])
        changed = False
        for tag, field in SEMANTIC_TAG_VALUES.items():
            if root.xpath(f"//w:sdt[w:sdtPr/w:tag/@w:val='{tag}']", namespaces=NS):
                _set_sdt_text(root, tag, scalar_values.get(field))
                changed = True
        if changed:
            parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def _render_cover(table: etree._Element, context: dict[str, Any]) -> None:
    scalars = context["scalar_slot_values"]
    if not scalars.get("has_separate_client"):
        return
    rows = list(table.xpath("./w:tr", namespaces=NS))
    if len(rows) != 3:
        raise ReportDomainError("TEMPLATE_COVER_TABLE_INVALID", "首页信息表结构异常。", status_code=500)
    client = _clone_row(rows[0], ["委托单位：", scalars["effective_client_name"]])
    rows[0].addnext(client)


def _checkbox(selected: bool) -> str:
    return "☒" if selected else "☐"


def _render_basic_information(table: etree._Element, context: dict[str, Any]) -> None:
    r2 = context["r2_context"]
    scalars = context["scalar_slot_values"]
    profile, metadata, phases = r2["profile"], r2["metadata"], r2["phases"]
    assessed = next((item for item in r2["organizations"] if item["organization_type"] == "assessed"), {})
    rows = list(table.xpath("./w:tr", namespaces=NS))
    if len(rows) != 37:
        raise ReportDomainError("TEMPLATE_BASIC_TABLE_INVALID", "基本信息表结构异常。", status_code=500)

    values = {
        1: [None, assessed.get("name", "")],
        2: [None, assessed.get("address", ""), None, assessed.get("postal_code", "")],
        3: [None, str(metadata.get("extension", {}).get("password_authority_department") or "")],
        4: [None, None, assessed.get("contact_name", ""), None, assessed.get("contact_title", "")],
        5: [None, None, assessed.get("contact_department", ""), None, assessed.get("office_phone", "")],
        6: [None, None, assessed.get("mobile_phone", ""), None, assessed.get("email", "")],
    }
    for row_index, row_values in values.items():
        cells = _cells(rows[row_index])
        for index, value in enumerate(row_values):
            if value is not None and index < len(cells):
                _set_cell_text(cells[index], value)

    critical = str(profile.get("critical_infrastructure_status") or "")
    _set_cell_text(
        _cells(rows[9])[1],
        f"{_checkbox(critical == 'recognized')}已认定，所属安全保护工作部门：{profile.get('critical_infrastructure_department','')}\n"
        f"{_checkbox(critical == 'not_recognized')}未认定",
    )
    filing = str(profile.get("level_filing_status") or "")
    filing_consistent = str(profile.get("level_filing_consistent") or "")
    _set_cell_text(
        _cells(rows[10])[1],
        f"{_checkbox(filing == 'filed')}已定级备案，第{metadata.get('classification_level') or '三级'}，"
        f"S {profile.get('level_filing_s','')}  A {profile.get('level_filing_a','')}  G {profile.get('level_filing_g','')}\n"
        f"备案证明编号：{profile.get('level_filing_number','')}\n"
        f"本次被测信息系统与等级保护定级系统是否一致：{_checkbox(filing_consistent == 'same')}是  "
        f"{_checkbox(filing_consistent == 'different')}否，变化情况说明：{profile.get('level_filing_difference','')}",
    )
    _set_cell_text(
        _cells(rows[11])[1],
        f"{_checkbox(filing == 'not_filed')}未定级，本次密评依据GB/T 39786—2021《信息安全技术 信息系统密码应用基本要求》"
        f"第{metadata.get('classification_level') or '三级'}信息系统要求",
    )
    assessment = str(profile.get("level_assessment_status") or "")
    assessment_texts = [
        f"{_checkbox(assessment == 'assessed')}已测评\n测评机构名称：{profile.get('level_assessment_organization','')}\n"
        f"测评时间：{profile.get('level_assessment_period','')}\n测评结论：{profile.get('level_assessment_conclusion','')}",
        f"{_checkbox(assessment == 'assessing')}正在测评\n测评机构名称：{profile.get('level_assessment_organization','')}",
        f"{_checkbox(assessment == 'not_assessed')}未测评",
    ]
    for cell, text in zip(_cells(rows[12])[1:], assessment_texts):
        _set_cell_text(cell, text)
    service_scope = profile.get("service_scope", {})
    scope = str(service_scope.get("kind") or "")
    scope_count = service_scope.get("count") or ""
    _set_cell_text(
        _cells(rows[13])[2],
        f"{_checkbox(scope == 'national')}全国  {_checkbox(scope == 'cross_province')}跨省（区、市）跨{scope_count}个  "
        f"{_checkbox(scope == 'province')}全省（区、市）  {_checkbox(scope == 'cross_city')}跨地（市、区）跨{scope_count}个  "
        f"{_checkbox(scope == 'local')}地（市、区）内  {_checkbox(scope == 'other')}其他 {service_scope.get('other','')}",
    )
    platform = profile.get("platform", {})
    operation = profile.get("operation", {})
    interconnection = profile.get("interconnection", {})
    cloud = profile.get("cloud_platform", {})
    plan = profile.get("crypto_plan", {})
    _set_cell_text(_cells(rows[16])[2], "  ".join(f"{_checkbox(value in platform.get('coverage', []))}{value}" for value in ("局域网", "城域网", "广域网", "其他")))
    _set_cell_text(_cells(rows[17])[2], "  ".join(f"{_checkbox(value in platform.get('nature', []))}{value}" for value in ("业务专网", "互联网", "其他")))
    _set_cell_text(_cells(rows[18])[1], str(platform.get("user_count") or operation.get("construction_stage") or ""))
    running = str(operation.get("status") or "")
    _set_cell_text(_cells(rows[19])[1], f"{_checkbox(running == 'running')}是，投入运行时间：{operation.get('started_at','')}  {_checkbox(running == 'not_running')}否，目前情况：{operation.get('construction_stage','')}")
    connection_types = list(interconnection.get("types") or [])
    _set_cell_text(_cells(rows[20])[1], "  ".join(f"{_checkbox(value in connection_types)}{value}" for value in ("与其他行业系统连接", "与本行业其他单位系统连接", "与本单位其他系统连接", "其他")) + f"\n互联系统名称：{interconnection.get('system_names','')}")
    dependency = str(cloud.get("dependency") or "")
    cloud_status = str(cloud.get("assessment_status") or "")
    _set_cell_text(_cells(rows[21])[1], f"{_checkbox(dependency == 'yes')}是，云平台名称：{cloud.get('name','')}")
    _set_cell_text(_cells(rows[21])[2], f"{_checkbox(cloud_status == 'assessed')}云平台已评估  {_checkbox(cloud_status == 'assessing')}云平台正在评估  {_checkbox(cloud_status == 'not_assessed')}云平台未评估\n密评机构名称：{cloud.get('organization','')}\n评估时间：{cloud.get('date','')}  评估结论：{cloud.get('conclusion','')}")
    _set_cell_text(_cells(rows[22])[1], f"{_checkbox(dependency == 'no')}否")
    plan_status = str(plan.get("status") or "")
    mode = str(plan.get("mode") or "")
    _set_cell_text(_cells(rows[23])[1], f"{_checkbox(plan_status == 'passed')}有密码应用方案，且通过密评，通过时间：{plan.get('passed_at','')}\n密评方式：{_checkbox(mode == 'self')}自行评估  {_checkbox(mode == 'commissioned')}委托密评机构评估，密评机构名称：{plan.get('organization','')}")
    _set_cell_text(_cells(rows[24])[1], f"{_checkbox(plan_status == 'not_passed')}有密码应用方案，但未通过密评")
    _set_cell_text(_cells(rows[25])[1], f"{_checkbox(plan_status == 'none')}无密码应用方案")
    products = r2["products"]
    total = sum(int(item.get("normalized_quantity") or 0) for item in products)
    exclusive = sum(int(item.get("normalized_quantity") or 0) for item in products if item.get("use_mode") == "exclusive")
    shared = sum(int(item.get("normalized_quantity") or 0) for item in products if item.get("use_mode") == "shared")
    counts = {kind: sum(int(item.get("normalized_quantity") or 0) for item in products if item.get("classification") == kind) for kind in ("certified", "uncertified_domestic", "foreign")}
    no_products = bool(profile.get("no_crypto_products"))
    _set_cell_text(_cells(rows[26])[1], f"{_checkbox(not no_products)}系统使用的密码产品{total}（台/套），独立使用{exclusive}（台/套），共享使用{shared}（台/套）；其中，取得认证证书的产品数量{counts['certified']}台/套，未取得认证证书的国内产品数量{counts['uncertified_domestic']}（台/套），国外产品数量{counts['foreign']}（台/套）。\n{_checkbox(no_products)}系统未使用密码产品")
    algorithms = [str(value) for value in profile.get("selected_algorithms", [])]
    catalog = ("SM1", "SM4", "SM7", "AES", "DES", "3DES", "SM2", "SM9", "RSA1024", "RSA2048", "SM3", "SHA-1", "SHA-256", "SHA-384", "SHA-512", "MD5", "ZUC")
    _set_cell_text(_cells(rows[27])[1], "  ".join(f"{_checkbox(value in algorithms)}{value}" for value in catalog) + "\n其他算法：" + "、".join(value for value in algorithms if value not in catalog))

    members = {str(item["member_uuid"]): item for item in r2["members"]}
    approval_rows = (
        (34, "compiler_member_uuid", phases.get("analysis_end")),
        (35, "reviewer_member_uuid", phases.get("report_review_at")),
        (36, "approver_member_uuid", phases.get("approved_at")),
    )
    for row_index, member_field, date in approval_rows:
        cells = _cells(rows[row_index])
        member = members.get(str(metadata.get(member_field) or ""), {})
        if len(cells) >= 4:
            _set_cell_text(cells[2], member.get("name", ""))
            _set_cell_text(cells[4] if len(cells) > 4 else cells[-1], str(date or ""))


def _render_assessment_conclusion_table(table: etree._Element, context: dict[str, Any]) -> None:
    blocks = context["chapter_blocks"]
    scalars = context["scalar_slot_values"]
    rows = list(table.xpath("./w:tr", namespaces=NS))
    _set_cell_text(_cells(rows[1])[1], blocks.get("conclusion.system_summary", {}).get("text") or scalars.get("system_overview") or "/")
    _set_cell_text(_cells(rows[2])[1], blocks.get("conclusion.assessment_summary", {}).get("text") or "/")
    cells = _cells(rows[3])
    if len(cells) >= 4:
        _set_cell_text(cells[1], scalars.get("conclusion") or "/")
        _set_cell_text(cells[3], scalars.get("overall_score") or "/")
    cells = _cells(rows[4])
    if len(cells) >= 4:
        _set_cell_text(cells[1], f"{scalars.get('not_applicable_count', 0)}/{scalars.get('indicator_total', 0)}")
        _set_cell_text(cells[3], scalars.get("high_risk_count", 0))


def _render_overall_evaluation(root: etree._Element, context: dict[str, Any]) -> None:
    blocks = context["chapter_blocks"]
    intro = _find_paragraph(
        root,
        startswith="本次信息系统商用密码应用安全性评估依据GB/T 39786—2021",
        contains="选取的测评指标总数",
    )
    _set_paragraph_text(intro, blocks.get("overall_evaluation.intro", {}).get("text") or "/", italic=False)
    for index, layer_name in enumerate(LAYER_NAMES, start=1):
        block = blocks.get(f"overall_evaluation.layer.{index}", {})
        first = _find_paragraph(root, startswith=f"在{layer_name}方面，")
        second = first.getnext()
        while second is not None and second.tag != Q("p"):
            second = second.getnext()
        situation = str(block.get("situation_description") or "未形成可用测评结果")
        stats = block.get("statistics") or {}
        _set_paragraph_text(first, f"{index}. 在{layer_name}方面，{situation}。", italic=False)
        _set_paragraph_text(
            second,
            f"测评结果：符合项{stats.get('compliant', 0)}项，部分符合项{stats.get('partially_compliant', 0)}项，"
            f"不符合项{stats.get('noncompliant', 0)}项，不适用项{stats.get('not_applicable', 0)}项。",
            italic=False,
        )
    outro = _find_paragraph(root, startswith="通过对【被测系统】的物理和环境安全")
    _set_paragraph_text(outro, blocks.get("overall_evaluation.outro", {}).get("text") or "/", italic=False)


def _clear_paragraphs(paragraphs: list[etree._Element]) -> None:
    for paragraph in paragraphs:
        _set_paragraph_text(paragraph, "", italic=False)


def _insert_paragraph_before(reference: etree._Element, template: etree._Element, text: str) -> None:
    paragraph = copy.deepcopy(template)
    _set_paragraph_text(paragraph, text, italic=False)
    reference.addprevious(paragraph)


def _render_security_issues(root: etree._Element, context: dict[str, Any]) -> None:
    blocks = context["chapter_blocks"]
    all_paragraphs = _paragraphs(root)
    security_title = _find_paragraph(root, exact="安全问题及改进建议")
    chapter_one = _find_paragraph(root, exact="测评项目概述")
    security_range = all_paragraphs[
        all_paragraphs.index(security_title) + 1:all_paragraphs.index(chapter_one)
    ]
    candidates = [
        p for p in security_range
        if _text(p).startswith("本次信息系统商用密码应用安全性评估依据GB/T 39786—2021")
        and "安全问题" in _text(p)
    ]
    if len(candidates) != 1:
        raise ReportDomainError("TEMPLATE_NARRATIVE_ANCHOR_INVALID", "安全问题引言锚点异常。", status_code=500)
    intro = candidates[0]
    intro_text = blocks.get("security_issues.intro", {}).get("text") or ""
    recommendation_intro = blocks.get("recommendations.intro", {}).get("text") or ""
    _set_paragraph_text(intro, f"{intro_text}{recommendation_intro}" or "/", italic=False)
    layer_headings = []
    for name in LAYER_NAMES:
        matches = [paragraph for paragraph in security_range if _text(paragraph) == name]
        if len(matches) != 1:
            raise ReportDomainError(
                "TEMPLATE_NARRATIVE_ANCHOR_INVALID", "安全问题层面标题锚点异常。",
                status_code=500, details={"layer": name, "count": len(matches)},
            )
        layer_headings.append(matches[0])
    for index, (layer_name, heading) in enumerate(zip(LAYER_NAMES, layer_headings), start=1):
        next_heading = layer_headings[index] if index < len(layer_headings) else chapter_one
        all_paragraphs = _paragraphs(root)
        start_index = all_paragraphs.index(heading)
        end_index = all_paragraphs.index(next_heading)
        segment = all_paragraphs[start_index:end_index]
        problem_label = next((p for p in segment if _text(p) == "问题描述："), None)
        recommendation_label = next((p for p in segment if _text(p) == "改进建议："), None)
        layer = blocks.get(f"security_issues.layer.{index}", {})
        recommendations = blocks.get(f"recommendations.layer.{index}", {})
        if not layer.get("visible"):
            _clear_paragraphs(segment)
            continue
        if problem_label is None or recommendation_label is None:
            raise ReportDomainError("TEMPLATE_NARRATIVE_ANCHOR_INVALID", "安全问题层面槽位异常。", status_code=500, details={"layer": layer_name})
        _set_paragraph_text(heading, f"{index}. {layer_name}", italic=False)
        between_problem = segment[segment.index(problem_label) + 1:segment.index(recommendation_label)]
        _clear_paragraphs(between_problem)
        problem_template = between_problem[0] if between_problem else problem_label
        for item in layer.get("problems", []):
            _insert_paragraph_before(
                recommendation_label,
                problem_template,
                f"{item.get('number')}）{item.get('indicator_name')}：{item.get('description')}",
            )
        all_paragraphs = _paragraphs(root)
        next_index = all_paragraphs.index(next_heading)
        recommendation_index = all_paragraphs.index(recommendation_label)
        after_recommendation = all_paragraphs[recommendation_index + 1:next_index]
        _clear_paragraphs(after_recommendation)
        recommendation_template = after_recommendation[0] if after_recommendation else recommendation_label
        for item in recommendations.get("items", []):
            _insert_paragraph_before(next_heading, recommendation_template, f"{item.get('number')}）{item.get('text')}")


def _render_narratives(root: etree._Element, context: dict[str, Any]) -> None:
    blocks = context["chapter_blocks"]
    scalars = context["scalar_slot_values"]
    _render_overall_evaluation(root, context)
    _render_security_issues(root, context)
    objective = _find_paragraph(root, startswith="中互金认证有限公司受【被测单位】委托")
    _set_paragraph_text(
        objective,
        f"中互金认证有限公司受{scalars.get('effective_client_name') or '/'}委托，于{scalars.get('assessment_start') or '/'}至"
        f"{scalars.get('assessment_end') or '/'}，依据GB/T 39786—2021《信息安全技术 信息系统密码应用基本要求》的"
        f"第{scalars.get('security_level') or '三级'}别要求，对{scalars.get('assessed_name') or '/'}的"
        f"{scalars.get('system_name') or '/'}从物理和环境安全、网络和通信安全、设备和计算安全、应用和数据安全、管理制度、"
        "人员管理、建设运行和应急处置等方面进行商用密码应用安全性评估，通过测评项目的实施，根据被测信息系统当前的安全状况，"
        "给出测评结果并提出改进建议，以确保被测信息系统达到GB/T 39786—2021《信息安全技术 信息系统密码应用基本要求》的要求，"
        "也为其信息资产安全和业务持续稳定运行提供保障。",
        italic=False,
    )
    risk = _find_paragraph(root, startswith="根据《商用密码应用安全性评估高风险判定指引》【高风险判定】")
    _set_paragraph_text(risk, blocks.get("risk_analysis.summary", {}).get("summary_text") or "/", italic=False)
    chapter7 = _find_paragraph(root, startswith="通过对【被测单位】的【被测系统】的物理和环境安全")
    _set_paragraph_text(chapter7, blocks.get("assessment_conclusion", {}).get("text") or "/", italic=False)
    for paragraph in _paragraphs(root):
        if (
            re.fullmatch(r"图A-[1-8]-\d+\s+[xX]{2,}", _text(paragraph))
            or _text(paragraph).startswith("测评记录中若存在图片")
        ):
            _set_paragraph_text(paragraph, "", italic=False)


def _generic_values(number: int, row: dict[str, Any]) -> list[Any]:
    if number in {20, 21, 22, 23, 24}:
        methods = set(row.get("methods") or [])
        method_text = "".join(f"{_checkbox(name in methods)}{name}" for name in ("访谈", "文档审查", "现场检查", "配置检查", "工具测试"))
        if number == 24:
            return [row.get("number"), row.get("unit"), row.get("object_name"), method_text, row.get("remark")]
        return [row.get("number"), row.get("object_name"), method_text, row.get("remark")]
    expected_columns = {4: 4, 5: 4, 6: 5, 7: 7, 8: 9, 9: 7, 10: 6, 11: 5, 12: 6, 13: 3, 14: 5, 15: 3, 19: 5, 53: 4}
    raw_values = row.get("__values__")
    if isinstance(raw_values, list):
        expected = expected_columns.get(number, len(raw_values))
        if len(raw_values) == expected:
            return list(raw_values)
        if len(raw_values) == expected - 1:
            return [row.get("number", "")] + list(raw_values)
    key_orders = {
        4: ("number", "name", "role", "qualification"),
        5: ("number", "name", "address", "remark"),
        6: ("number", "name", "type", "deployment", "remark"),
        7: ("number", "product_name", "manufacturer_model", "certificate_number", "algorithms", "quantity_text", "purpose"),
        8: ("number", "name", "manufacturer", "model", "operating_system", "virtual", "purpose", "quantity", "importance"),
        9: ("number", "name", "manufacturer", "model", "purpose", "quantity", "importance"),
        10: ("number", "name", "version", "location", "description", "importance"),
        11: ("number", "name", "version", "location", "description"),
        12: ("number", "name", "description", "application", "storage", "security_requirement"),
        13: ("number", "name", "description"),
        14: ("number", "name", "role", "responsibilities", "contact"),
        15: ("number", "name", "provider"),
        19: ("number", "name", "version", "purpose", "remark"),
    }
    order = key_orders.get(number)
    if order is None:
        return list(row.values())
    return [row.get(key, "") for key in order]


def _render_generic_tables(root: etree._Element, context: dict[str, Any]) -> None:
    row_map = context["table_rows_by_table_id"]
    for number in (*range(4, 16), 19, 20, 21, 22, 23, 24):
        table = _table_by_anchor(root, number)
        source = row_map.get(f"report_table_{number:03d}", [])
        prepared = []
        for index, row in enumerate(source, start=1):
            item = dict(row)
            item.setdefault("number", index)
            prepared.append(_generic_values(number, item))
        _replace_rows(table, header_rows=1, source_rows=prepared)

    table16 = _table_by_anchor(root, 16)
    threat_rows = [[item.get("number"), item.get("threat"), item.get("frequency") or "/"] for item in row_map.get("report_table_016", [])]
    _replace_rows(table16, header_rows=1, source_rows=threat_rows)


def _render_special_indicators(root: etree._Element, context: dict[str, Any]) -> None:
    special = context["r2_context"].get("special_indicators", [])
    if not special:
        return
    standards = {
        str(item.get("standard_uuid") or ""): item
        for item in context["r2_context"].get("standards", [])
    }
    table = _table_by_anchor(root, 17)
    rows = list(table.xpath("./w:tr", namespaces=NS))
    if len(rows) < 3:
        raise ReportDomainError(
            "TEMPLATE_SPECIAL_INDICATOR_TABLE_INVALID",
            "特殊指标表结构异常。",
            status_code=500,
        )
    template_row = rows[1]
    summary = rows[-1]
    for item in special:
        standard = standards.get(str(item.get("manual_standard_uuid") or ""), {})
        reference = " ".join(
            value for value in (
                str(standard.get("standard_code") or "").strip(),
                str(standard.get("standard_name") or "").strip(),
            ) if value
        )
        summary.addprevious(
            _clone_row(
                template_row,
                [
                    "特殊指标", "", item.get("indicator_name") or item.get("indicator_code") or "/",
                    item.get("description") or "/", reference or "/",
                ],
            )
        )
    summary_cells = _cells(summary)
    if summary_cells:
        _set_cell_text(summary_cells[-1], f"41项（另含特殊指标{len(special)}项）")


def _render_object_tables(root: etree._Element, context: dict[str, Any]) -> None:
    row_map = context["table_rows_by_table_id"]
    for number in (20, 21, 22, 23):
        rows = row_map.get(f"report_table_{number:03d}", [])
        prepared = []
        for index, row in enumerate(rows, start=1):
            item = dict(row)
            item["number"] = index
            prepared.append(_generic_values(number, item))
        _replace_rows(_table_by_anchor(root, number), header_rows=1, source_rows=prepared)


def _render_not_applicable_table(root: etree._Element, context: dict[str, Any]) -> None:
    final = context["r3_context"].get("final_projection", {})
    items = [item for item in final.get("indicators", []) if item.get("indicator_result") == "不适用"]
    rows = [[index, item.get("indicator_name"), "该指标全部测评对象均判定为不适用。", "附录A"] for index, item in enumerate(items, start=1)]
    _replace_rows(_table_by_anchor(root, 18), header_rows=1, source_rows=rows)


def _render_chapter4(root: etree._Element, context: dict[str, Any]) -> None:
    projections = context["r3_context"].get("original_projection", {}).get("chapter4_tables", {})
    for offset, number in enumerate(range(25, 36), start=1):
        table = _table_by_anchor(root, number)
        projection = projections.get(f"table_4_{offset}", {})
        columns = list(projection.get("columns") or [])
        body = [
            [index, row.get("object_name") or "/"]
            + [row.get("cells", {}).get(column, {}).get("result") or "不适用" for column in columns]
            for index, row in enumerate(projection.get("rows") or [], start=1)
        ]
        summary = ["单元测评结果（符合/部分符合/不符合/不适用）"] + [
            projection.get("summary", {}).get(column, {}).get("result") or "不适用"
            for column in columns
        ]
        header_rows = 1 if number in {29, 30, 31} else 2
        _replace_rows(table, header_rows=header_rows, source_rows=body, keep_summary=True, summary_values=summary)


def _render_corrections(root: etree._Element, context: dict[str, Any]) -> None:
    corrections = context["r3_context"].get("correction_projection", {}).get("rows", [])
    final_rows = context["r3_context"].get("final_projection", {}).get("rows", [])
    by_id = {int(row["source_row_id"]): row for row in final_rows if row.get("source_row_id") is not None}
    layer_names = {
        "A-1": "物理和环境安全", "A-2": "网络和通信安全",
        "A-3": "设备和计算安全", "A-4": "应用和数据安全",
    }
    rows = []
    for index, correction in enumerate(corrections, start=1):
        related = [by_id.get(int(source_id), {}) for source_id in correction.get("related_source_row_ids", [])]
        rows.append(
            [
                index,
                layer_names.get(str(correction.get("section_code") or ""), "/"),
                correction.get("indicator_name"),
                correction.get("object_name"),
                f"{correction.get('original_result') or '/'} / {correction.get('original_score') or '/'}",
                f"{correction.get('final_result') or '/'} / {correction.get('final_score') or '/'}",
                (
                    "关联测评对象："
                    + ("、".join(str(item.get("object_name") or item.get("object_uuid") or "") for item in related) or "/")
                ),
            ]
        )
    if not rows:
        rows = [["/"] * 7]
    _replace_rows(_table_by_anchor(root, 36), header_rows=1, source_rows=rows)


def _render_overall_score(root: etree._Element, context: dict[str, Any]) -> None:
    table = _table_by_anchor(root, 37)
    rows = list(table.xpath("./w:tr", namespaces=NS))
    indicators = context["r3_context"].get("final_projection", {}).get("indicators", [])
    score = context["r3_context"].get("final_projection", {}).get("score", {})
    layer_scores = {str(item.get("layer_code")): item.get("score") for item in score.get("layers", [])}
    seen_layers: set[str] = set()
    result_index = {"符合": 2, "部分符合": 3, "不符合": 4, "不适用": 5}
    for offset, indicator in enumerate(indicators, start=2):
        if offset >= len(rows) - 2:
            break
        cells = _cells(rows[offset])
        selected = result_index.get(str(indicator.get("indicator_result")))
        for index in range(2, min(6, len(cells))):
            _set_cell_text(cells[index], "1" if index == selected else "0")
        if len(cells) > 6:
            _set_cell_text(cells[6], indicator.get("unit_score") or "/")
        layer = str(indicator.get("layer_code") or "")
        if len(cells) > 7 and layer not in seen_layers:
            _set_cell_text(cells[7], layer_scores.get(layer) or "/")
            seen_layers.add(layer)
    total = context["r3_context"].get("final_projection", {}).get("statistics", {}).get("total", {})
    total_row = _cells(rows[-2])
    values = [total.get("compliant", 0), total.get("partially_compliant", 0), total.get("noncompliant", 0), total.get("not_applicable", 0), score.get("display_score") or "/"]
    for cell, value in zip(total_row[1:], values):
        _set_cell_text(cell, value)
    conclusion_row = _cells(rows[-1])
    if conclusion_row:
        _set_cell_text(conclusion_row[-1], context["scalar_slot_values"].get("conclusion") or "/")


def _render_risks(root: etree._Element, context: dict[str, Any]) -> None:
    block = context["chapter_blocks"].get("risk_analysis.rows", {})
    rows = []
    for item in block.get("rows", []):
        level = str(item.get("risk_level") or "")
        rows.append(
            [
                item.get("number"), item.get("layer_name"),
                f"{item.get('indicator_name')}：{item.get('problem_description')}",
                "、".join(item.get("threat_ids") or []), item.get("analysis"),
                "1" if level == "high" else "0", "1" if level == "medium" else "0", "1" if level == "low" else "0",
            ]
        )
    table = _table_by_anchor(root, 38)
    existing = list(table.xpath("./w:tr", namespaces=NS))
    data_template = existing[2]
    statistics_template = existing[-1]
    for row in existing[2:]:
        table.remove(row)
    for values in rows or [["/"] * 8]:
        table.append(_clone_row(data_template, values))
    stats = block.get("statistics", {})
    table.append(_clone_row(statistics_template, ["统计", "", "", "", "", stats.get("high", 0), stats.get("medium", 0), stats.get("low", 0)]))


def _appendix_content(
    context: dict[str, Any], section_code: str
) -> tuple[etree._Element, list[etree._Element], dict[str, bytes]]:
    profile = load_template_profile()
    images = [
        item for item in context["r2_context"].get("evidence", [])
        if item.get("section_code") == section_code
    ]
    figure_refs = build_figure_refs(section_code, images)
    rows = []
    for source in context["appendix_a_final_projection"].get("rows", []):
        if source.get("section_code") != section_code:
            continue
        corrected = bool(source.get("was_corrected"))
        score = str(source.get("final_object_score") or source.get("object_score") or "")
        rows.append(
            {
                "unit": source.get("indicator_name") or "",
                "object_name": source.get("object_name") or "",
                "record_text": source.get("record_text") or "",
                "d": "/" if corrected and section_code in {"A-1", "A-2", "A-3", "A-4"} else source.get("d"),
                "a": "/" if corrected and section_code in {"A-1", "A-2", "A-3", "A-4"} else source.get("a"),
                "k": "/" if corrected and section_code in {"A-1", "A-2", "A-3", "A-4"} else source.get("k"),
                "object_score": f"{score}*" if corrected and score != "/" else score,
                "unit_score": source.get("unit_score") or "",
                "compliance": source.get("object_result") or "不适用",
            }
        )
    document = Document()
    table = add_assessment_table(
        document,
        {"code": section_code},
        rows,
        profile,
        "final",
        figure_refs,
        authoritative_scores=True,
    )
    add_section_images(
        document,
        section_code,
        images,
        profile,
        BookmarkWriter(),
        figure_refs,
    )
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as package:
        temporary_parts = {name: package.read(name) for name in package.namelist()}
    temporary_document = etree.fromstring(temporary_parts["word/document.xml"])
    body_children = [
        copy.deepcopy(child)
        for child in temporary_document.xpath("/w:document/w:body/*[not(self::w:sectPr)]", namespaces=NS)
    ]
    if not body_children or body_children[0].tag != Q("tbl"):
        raise ReportDomainError(
            "APPENDIX_A_EMBED_RENDER_INVALID",
            "附录 A 嵌入式渲染未生成预期表格。",
            status_code=500,
            details={"section_code": section_code},
        )
    return body_children[0], body_children[1:], temporary_parts


def _next_relationship_id(relationships: etree._Element) -> int:
    numbers = []
    for value in relationships.xpath("/pr:Relationships/pr:Relationship/@Id", namespaces=NS):
        match = re.fullmatch(r"rId(\d+)", str(value))
        if match:
            numbers.append(int(match.group(1)))
    return max(numbers, default=0) + 1


def _import_appendix_resources(
    root: etree._Element,
    elements: list[etree._Element],
    temporary_parts: dict[str, bytes],
    parts: dict[str, bytes],
) -> None:
    if not elements:
        return
    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    temporary_relationships = etree.fromstring(temporary_parts["word/_rels/document.xml.rels"])
    next_relationship = _next_relationship_id(relationships)
    relationship_map: dict[str, str] = {}
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    temporary_content_types = etree.fromstring(temporary_parts["[Content_Types].xml"])
    content_type_namespace = etree.QName(content_types).namespace
    existing_extensions = {
        str(node.get("Extension") or "").lower()
        for node in content_types.findall(f"{{{content_type_namespace}}}Default")
    }
    temporary_defaults = {
        str(node.get("Extension") or "").lower(): str(node.get("ContentType") or "")
        for node in temporary_content_types.findall(f"{{{content_type_namespace}}}Default")
    }

    for relationship in temporary_relationships.xpath(
        "/pr:Relationships/pr:Relationship[contains(@Type, '/image')]", namespaces=NS
    ):
        old_id = str(relationship.get("Id") or "")
        target = str(relationship.get("Target") or "")
        source_name = f"word/{target}" if not target.startswith("/") else target.lstrip("/")
        data = temporary_parts.get(source_name)
        if not old_id or data is None:
            raise ReportDomainError(
                "APPENDIX_A_IMAGE_RELATION_INVALID",
                "附录 A 图片关系缺少目标资源。",
                status_code=500,
            )
        extension = Path(source_name).suffix.lower() or ".bin"
        digest = hashlib.sha256(data).hexdigest()
        target_name = f"media/r4_{digest[:24]}{extension}"
        package_name = f"word/{target_name}"
        if package_name in parts and parts[package_name] != data:
            raise ReportDomainError(
                "APPENDIX_A_IMAGE_NAME_CONFLICT",
                "附录 A 图片资源名称发生冲突。",
                status_code=500,
            )
        parts[package_name] = data
        new_id = f"rId{next_relationship}"
        next_relationship += 1
        relationship_map[old_id] = new_id
        imported = etree.SubElement(relationships, f"{{{REL}}}Relationship")
        imported.set("Id", new_id)
        imported.set("Type", str(relationship.get("Type")))
        imported.set("Target", target_name)
        suffix = extension.lstrip(".").lower()
        if suffix not in existing_extensions:
            default = etree.SubElement(content_types, f"{{{content_type_namespace}}}Default")
            default.set("Extension", suffix)
            default.set(
                "ContentType",
                temporary_defaults.get(suffix)
                or mimetypes.guess_type(f"file.{suffix}")[0]
                or "application/octet-stream",
            )
            existing_extensions.add(suffix)

    for element in elements:
        for node in element.xpath(".//*[@r:embed]", namespaces=NS):
            old_id = node.get(f"{{{R}}}embed")
            if old_id in relationship_map:
                node.set(f"{{{R}}}embed", relationship_map[old_id])

    current_doc_pr = [
        int(value) for value in root.xpath("//*[local-name()='docPr']/@id")
        if str(value).isdigit()
    ]
    next_doc_pr = max(current_doc_pr, default=0) + 1
    for element in elements:
        for node in element.xpath(".//*[local-name()='docPr']"):
            node.set("id", str(next_doc_pr))
            next_doc_pr += 1

    current_bookmark_ids = [
        int(value) for value in root.xpath("//w:bookmarkStart/@w:id", namespaces=NS)
        if str(value).lstrip("-").isdigit()
    ]
    next_bookmark = max(current_bookmark_ids, default=0) + 1
    bookmark_map: dict[str, str] = {}
    existing_names = set(root.xpath("//w:bookmarkStart/@w:name", namespaces=NS))
    for element in elements:
        for start in element.xpath(".//w:bookmarkStart", namespaces=NS):
            old_id = str(start.get(Q("id")) or "")
            name = str(start.get(Q("name")) or "")
            if name and name in existing_names:
                raise ReportDomainError(
                    "APPENDIX_A_BOOKMARK_CONFLICT",
                    "附录 A 图片书签名称发生冲突。",
                    status_code=500,
                    details={"bookmark": name},
                )
            new_id = str(next_bookmark)
            next_bookmark += 1
            bookmark_map[old_id] = new_id
            start.set(Q("id"), new_id)
            if name:
                existing_names.add(name)
        for end in element.xpath(".//w:bookmarkEnd", namespaces=NS):
            old_id = str(end.get(Q("id")) or "")
            if old_id in bookmark_map:
                end.set(Q("id"), bookmark_map[old_id])

    parts["word/_rels/document.xml.rels"] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    parts["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )


def _render_appendix_a(
    root: etree._Element, context: dict[str, Any], parts: dict[str, bytes]
) -> None:
    for number, section_code in zip(range(39, 47), (f"A-{index}" for index in range(1, 9))):
        original = _table_by_anchor(root, number)
        replacement, following, temporary_parts = _appendix_content(context, section_code)
        starts = original.xpath(
            f".//w:bookmarkStart[@w:name='rt_table_{number:03d}']", namespaces=NS
        )
        anchor_id = starts[0].get(Q("id")) if starts else None
        anchors = starts + (
            original.xpath(f".//w:bookmarkEnd[@w:id='{anchor_id}']", namespaces=NS)
            if anchor_id is not None else []
        )
        first_paragraph = replacement.find(f".//{Q('p')}")
        if first_paragraph is not None:
            for anchor in reversed(anchors):
                first_paragraph.insert(0, copy.deepcopy(anchor))
        original.getparent().replace(original, replacement)
        _import_appendix_resources(root, following, temporary_parts, parts)
        cursor = replacement
        for element in following:
            cursor.addnext(element)
            cursor = element


def _set_update_fields(parts: dict[str, bytes]) -> None:
    settings = etree.fromstring(parts["word/settings.xml"])
    updates = settings.xpath("/w:settings/w:updateFields", namespaces=NS)
    update = updates[0] if updates else etree.SubElement(settings, Q("updateFields"))
    update.set(Q("val"), "true")
    parts["word/settings.xml"] = etree.tostring(settings, xml_declaration=True, encoding="UTF-8", standalone="yes")


def _add_draft_marker(root: etree._Element) -> None:
    title = _find_paragraph(root, exact="商用密码应用安全性评估报告")
    marker = copy.deepcopy(title)
    _set_paragraph_text(marker, "草稿—未完成复核", italic=False)
    title.addnext(marker)


def _canonical_nodes(data: bytes, expression: str) -> list[bytes]:
    root = etree.fromstring(data)
    return [etree.tostring(node, method="c14n", exclusive=True) for node in root.xpath(expression, namespaces=NS)]


def _relationship_signature(data: bytes) -> set[tuple[str, str, str, str]]:
    root = etree.fromstring(data)
    return {
        (
            str(node.get("Id") or ""),
            str(node.get("Type") or ""),
            str(node.get("Target") or ""),
            str(node.get("TargetMode") or ""),
        )
        for node in root.xpath("/pr:Relationships/pr:Relationship", namespaces=NS)
    }


def _content_type_signature(data: bytes) -> set[tuple[str, str, str]]:
    root = etree.fromstring(data)
    namespace = etree.QName(root).namespace
    values = {
        ("Default", str(node.get("Extension") or "").lower(), str(node.get("ContentType") or ""))
        for node in root.findall(f"{{{namespace}}}Default")
    }
    values.update(
        ("Override", str(node.get("PartName") or ""), str(node.get("ContentType") or ""))
        for node in root.findall(f"{{{namespace}}}Override")
    )
    return values


def _assert_template_mutation_allowlist(
    original_parts: dict[str, bytes], rendered_parts: dict[str, bytes]
) -> dict[str, Any]:
    semantic_stories: set[str] = set()
    for name, data in original_parts.items():
        if not re.fullmatch(r"word/(?:header|footer)\d+\.xml", name):
            continue
        story = etree.fromstring(data)
        tags = set(story.xpath("//w:sdtPr/w:tag/@w:val", namespaces=NS))
        if tags.intersection(SEMANTIC_TAG_VALUES):
            semantic_stories.add(name)

    allowed_mutations = {
        "word/document.xml",
        "word/settings.xml",
        "word/_rels/document.xml.rels",
        "[Content_Types].xml",
        *semantic_stories,
    }
    original_names = set(original_parts)
    rendered_names = set(rendered_parts)
    removed = sorted(original_names - rendered_names)
    unexpected_changes = sorted(
        name
        for name in original_names & rendered_names
        if original_parts[name] != rendered_parts[name] and name not in allowed_mutations
    )
    new_parts = sorted(rendered_names - original_names)
    invalid_new_parts = [
        name for name in new_parts
        if not re.fullmatch(r"word/media/r4_[0-9a-f]{24}\.[A-Za-z0-9]+", name)
    ]

    original_sections = _canonical_nodes(original_parts["word/document.xml"], "//w:sectPr")
    rendered_sections = _canonical_nodes(rendered_parts["word/document.xml"], "//w:sectPr")
    sections_preserved = original_sections == rendered_sections

    original_relationships = _relationship_signature(original_parts["word/_rels/document.xml.rels"])
    rendered_relationships = _relationship_signature(rendered_parts["word/_rels/document.xml.rels"])
    missing_relationships = sorted(original_relationships - rendered_relationships)
    invalid_relationships = sorted(
        item for item in rendered_relationships - original_relationships
        if not (
            item[1].endswith("/image")
            and re.fullmatch(r"media/r4_[0-9a-f]{24}\.[A-Za-z0-9]+", item[2])
            and not item[3]
        )
    )

    original_content_types = _content_type_signature(original_parts["[Content_Types].xml"])
    rendered_content_types = _content_type_signature(rendered_parts["[Content_Types].xml"])
    missing_content_types = sorted(original_content_types - rendered_content_types)
    if (
        removed or unexpected_changes or invalid_new_parts or not sections_preserved
        or missing_relationships or invalid_relationships or missing_content_types
    ):
        raise ReportDomainError(
            "REPORT_TEMPLATE_MUTATION_ALLOWLIST_VIOLATION",
            "完整报告装配修改了母版白名单之外的结构。",
            status_code=500,
            details={
                "removed_parts": removed,
                "unexpected_changed_parts": unexpected_changes,
                "invalid_new_parts": invalid_new_parts,
                "sections_preserved": sections_preserved,
                "missing_relationships": missing_relationships,
                "invalid_relationships": invalid_relationships,
                "missing_content_types": missing_content_types,
            },
        )
    return {
        "template_allowlist_verified": True,
        "preserved_section_signatures": len(rendered_sections),
        "unchanged_original_parts": sum(
            original_parts[name] == rendered_parts[name] for name in original_names & rendered_names
        ),
        "new_media_parts": len(new_parts),
    }


def render_report(context: dict[str, Any], destination: Path) -> dict[str, Any]:
    package = report_template_registry.load()
    with zipfile.ZipFile(io.BytesIO(package.runtime_template_bytes)) as source:
        infos = source.infolist()
        parts = {info.filename: source.read(info.filename) for info in infos}
    original_parts = dict(parts)
    root = etree.fromstring(parts["word/document.xml"])
    _render_cover(_table_by_anchor(root, 1), context)
    _render_basic_information(_table_by_anchor(root, 2), context)
    _render_assessment_conclusion_table(_table_by_anchor(root, 3), context)
    _render_narratives(root, context)
    _render_generic_tables(root, context)
    _render_special_indicators(root, context)
    _render_not_applicable_table(root, context)
    _render_chapter4(root, context)
    _render_corrections(root, context)
    _render_overall_score(root, context)
    _render_risks(root, context)
    _render_appendix_a(root, context, parts)
    if context["project_identity"]["export_mode"] == "draft":
        _add_draft_marker(root)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
    # Narrative anchors are defined against the frozen master text.  Populate
    # semantic SDTs only after those deterministic replacements so a scalar
    # value cannot invalidate a later master fingerprint.
    _render_semantic_slots(parts, context["scalar_slot_values"])
    _set_update_fields(parts)
    allowlist_result = _assert_template_mutation_allowlist(original_parts, parts)

    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as output:
        original_names = {info.filename for info in infos}
        for info in infos:
            output.writestr(info, parts[info.filename])
        for name in sorted(set(parts) - original_names):
            output.writestr(name, parts[name])
    validation = validate_rendered_report(destination, final=context["project_identity"]["export_mode"] == "final")
    return {**validation, **allowlist_result}


def validate_rendered_report(path: Path, *, final: bool) -> dict[str, Any]:
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        root = etree.fromstring(package.read("word/document.xml"))
        story_names = [
            name for name in names
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
        ]
        story_roots = [etree.fromstring(package.read(name)) for name in story_names]
        relationship_roots = [
            etree.fromstring(package.read(name)) for name in names if name.endswith(".rels")
        ]
    tables = root.xpath("//w:tbl", namespaces=NS)
    sections = root.xpath("//w:sectPr", namespaces=NS)
    if len(tables) != 55 or len(sections) != 17:
        raise ReportDomainError(
            "REPORT_DOCX_STRUCTURE_INVALID", "生成报告的分节或表格数量不符合母版。", status_code=500,
            details={"sections": len(sections), "tables": len(tables)},
        )
    forbidden_parts = [
        name for name in names
        if name.startswith(("word/activeX/", "word/embeddings/", "customXml/", "_xmlsignatures/"))
        or "comments" in name.lower() or name.endswith("vbaProject.bin")
    ]
    external = sum(len(root.xpath("/pr:Relationships/pr:Relationship[@TargetMode='External']", namespaces=NS)) for root in relationship_roots)
    revisions = sum(
        len(story.xpath("//w:ins | //w:del | //w:moveFrom | //w:moveTo", namespaces=NS))
        for story in story_roots
    )
    instructions = " ".join(
        instruction
        for story in story_roots
        for instruction in story.xpath("//w:instrText/text() | //w:fldSimple/@w:instr", namespaces=NS)
    )
    visible = "".join(
        text
        for story in story_roots
        for text in story.xpath("//w:t/text()", namespaces=NS)
    )
    tags = [
        tag
        for story in story_roots
        for tag in story.xpath("//w:sdtPr/w:tag/@w:val", namespaces=NS)
    ]
    private_factors = [value for value in tags if value.lower() in {"ra", "rk"}]
    if re.search(r"\b(?:Ra|Rk)\b", visible):
        private_factors.append("visible_text")
    active_italic_expression = (
        "//w:i[not(ancestor::m:oMath) and "
        "not(@w:val='0' or @w:val='false' or @w:val='off')] | "
        "//w:iCs[not(ancestor::m:oMath) and "
        "not(@w:val='0' or @w:val='false' or @w:val='off')]"
    )
    active_italics = sum(
        len(story.xpath(active_italic_expression, namespaces=NS))
        for story in story_roots
    )
    if (
        forbidden_parts or external or revisions
        or re.search(r"\bNUMPAGES\b", instructions, re.IGNORECASE)
        or private_factors or active_italics
    ):
        raise ReportDomainError(
            "REPORT_DOCX_FORBIDDEN_CONTENT", "生成报告包含禁止的 OOXML 内容。", status_code=500,
            details={
                "forbidden_parts": forbidden_parts, "external_relationships": external,
                "revisions": revisions, "private_factors": private_factors,
                "active_non_formula_italics": active_italics,
            },
        )
    placeholders = sorted(set(PLACEHOLDER_RE.findall(visible)))
    if final and placeholders:
        raise ReportDomainError(
            "REPORT_DOCX_PLACEHOLDER_REMAINS", "正式报告仍包含未替换占位符。", status_code=422,
            details={"placeholders": placeholders[:30]},
        )
    required_fields = {"TOC", "PAGE", "PAGEREF", "SEQ", "REF"}
    available_fields = {name for name in required_fields if re.search(rf"\b{name}\b", instructions, re.IGNORECASE)}
    if available_fields != required_fields:
        raise ReportDomainError(
            "REPORT_DOCX_FIELD_CONTRACT_INVALID", "目录、页码或交叉引用字段不完整。", status_code=500,
            details={"missing": sorted(required_fields - available_fields)},
        )
    bookmark_names = set(root.xpath("//w:bookmarkStart/@w:name", namespaces=NS))
    reference_targets = re.findall(r"\b(?:REF|PAGEREF)\s+([A-Za-z_][A-Za-z0-9_]*)", instructions, re.IGNORECASE)
    missing_targets = sorted({target for target in reference_targets if target not in bookmark_names})
    if missing_targets:
        raise ReportDomainError(
            "REPORT_DOCX_REFERENCE_TARGET_MISSING",
            "生成报告存在缺失的书签或交叉引用目标。",
            status_code=500,
            details={"missing_targets": missing_targets[:50]},
        )
    return {
        "section_count": len(sections),
        "table_count": len(tables),
        "field_types": sorted(available_fields),
        "placeholder_count": len(placeholders),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }
